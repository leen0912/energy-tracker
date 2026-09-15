from pathlib import Path
from datetime import date

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parent
OUT_DOCX = ROOT / "给梁姐的需求演化简报-生活时间显影器.docx"
OUT_PNG = ROOT / "需求演化地图-给梁姐.png"


INK = RGBColor(38, 42, 52)
MUTED = RGBColor(92, 98, 112)
BLUE = RGBColor(48, 91, 142)
TEAL = RGBColor(38, 128, 123)
GOLD = RGBColor(174, 123, 40)
ROSE = RGBColor(168, 76, 92)
LIGHT_BLUE = "EAF2FB"
LIGHT_TEAL = "E9F5F3"
LIGHT_GOLD = "FFF4DF"
LIGHT_ROSE = "FBECEF"
LIGHT_GRAY = "F5F6F8"


def font_path(*names):
    font_dirs = [
        Path("C:/Windows/Fonts"),
        Path("/usr/share/fonts"),
        Path("/usr/local/share/fonts"),
    ]
    for d in font_dirs:
        for name in names:
            p = d / name
            if p.exists():
                return str(p)
    return None


FONT_REG = font_path("msyh.ttc", "simsun.ttc", "NotoSansCJK-Regular.ttc")
FONT_BOLD = font_path("msyhbd.ttc", "simhei.ttf", "NotoSansCJK-Bold.ttc") or FONT_REG


def set_run_font(run, size=None, color=None, bold=None):
    name = "Microsoft YaHei"
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold


def set_para(p, before=0, after=6, line=1.25, align=None):
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line
    if align is not None:
        p.alignment = align


def add_p(doc, text="", size=10.8, color=INK, bold=False, after=6, before=0, align=None):
    p = doc.add_paragraph()
    set_para(p, before=before, after=after, align=align)
    r = p.add_run(text)
    set_run_font(r, size=size, color=color, bold=bold)
    return p


def add_heading(doc, text, level=1):
    size = 16 if level == 1 else 13 if level == 2 else 11.5
    color = BLUE if level <= 2 else TEAL
    before = 16 if level == 1 else 10
    after = 7 if level == 1 else 5
    p = doc.add_paragraph()
    set_para(p, before=before, after=after, line=1.15)
    r = p.add_run(text)
    set_run_font(r, size=size, color=color, bold=True)
    return p


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=140, bottom=100, end=140):
    tc_pr = cell._tc.get_or_add_tcPr()
    mar = tc_pr.first_child_found_in("w:tcMar")
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tc_pr.append(mar)
    for m, v in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="DADDE5"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        tag = f"w:{edge}"
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "6")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)


def table_width(table, widths):
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(int(width * 1440)))
            tc_w.set(qn("w:type"), "dxa")


def add_callout(doc, title, body, fill=LIGHT_BLUE, accent=BLUE):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table_width(table, [6.5])
    set_table_borders(table, color=fill)
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_margins(cell, top=170, bottom=170, start=220, end=220)
    p = cell.paragraphs[0]
    set_para(p, after=4, line=1.2)
    r = p.add_run(title)
    set_run_font(r, size=11, color=accent, bold=True)
    p2 = cell.add_paragraph()
    set_para(p2, after=0, line=1.25)
    r2 = p2.add_run(body)
    set_run_font(r2, size=10.5, color=INK)
    doc.add_paragraph()


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style=None)
        p.style = doc.styles["List Bullet"]
        set_para(p, after=4, line=1.2)
        r = p.add_run(item)
        set_run_font(r, size=10.6, color=INK)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style=None)
        p.style = doc.styles["List Number"]
        set_para(p, after=4, line=1.2)
        r = p.add_run(item)
        set_run_font(r, size=10.6, color=INK)


def draw_wrapped(draw, text, xy, font, fill, max_width, line_gap=8):
    x, y = xy
    lines = []
    buf = ""
    for ch in text:
        trial = buf + ch
        if draw.textlength(trial, font=font) <= max_width:
            buf = trial
        else:
            if buf:
                lines.append(buf)
            buf = ch
    if buf:
        lines.append(buf)
    for line in lines:
        draw.text((x, y), line, font=font, fill=fill)
        y += font.size + line_gap
    return y


def create_infographic():
    w, h = 1600, 1100
    img = Image.new("RGB", (w, h), "#F7F3EA")
    draw = ImageDraw.Draw(img)
    title_font = ImageFont.truetype(FONT_BOLD, 62) if FONT_BOLD else ImageFont.load_default()
    sub_font = ImageFont.truetype(FONT_REG, 30) if FONT_REG else ImageFont.load_default()
    node_font = ImageFont.truetype(FONT_BOLD, 30) if FONT_BOLD else ImageFont.load_default()
    body_font = ImageFont.truetype(FONT_REG, 24) if FONT_REG else ImageFont.load_default()
    small_font = ImageFont.truetype(FONT_REG, 22) if FONT_REG else ImageFont.load_default()

    draw.rounded_rectangle((56, 50, 1544, 1048), radius=42, fill="#FFFDF8", outline="#E0D8C8", width=2)
    draw.text((110, 105), "从番茄钟到生活时间显影器", font=title_font, fill="#26313A")
    draw.text((114, 184), "一次需求对话的演化：不是逼自己高效，而是看见自己真实地生活过。", font=sub_font, fill="#667085")

    nodes = [
        ("1", "起点", "付费番茄钟太烦，想要一个自己的工具", "#EAF2FB", "#305B8E"),
        ("2", "修正", "开始时很难知道要做多久，所以不要先设置时长", "#E9F5F3", "#26807B"),
        ("3", "再定义", "时钟应在后台归档，前台只负责帮我开始", "#FFF4DF", "#AE7B28"),
        ("4", "更深层", "低电量生活：睡眠、饮食、咖啡因与精力状态", "#FBECEF", "#A84C5C"),
        ("5", "真正动机", "漂亮总结图、热力图和翻日历的满足感", "#F0ECFF", "#6B5BA7"),
        ("6", "最终形态", "读取已有日历 + 聊感受 + 温柔提醒 + 生成可分享生活日历", "#EAF7ED", "#497A4A"),
    ]

    start_x, start_y = 120, 290
    card_w, card_h = 410, 188
    gap_x, gap_y = 58, 64
    positions = []
    for i in range(6):
        row, col = divmod(i, 3)
        x = start_x + col * (card_w + gap_x)
        y = start_y + row * (card_h + gap_y)
        positions.append((x, y))

    for i, ((num, heading, text, fill, accent), (x, y)) in enumerate(zip(nodes, positions)):
        draw.rounded_rectangle((x, y, x + card_w, y + card_h), radius=26, fill=fill, outline="#FFFFFF", width=3)
        draw.ellipse((x + 24, y + 24, x + 72, y + 72), fill=accent)
        draw.text((x + 41, y + 31), num, font=small_font, fill="#FFFFFF", anchor="mm")
        draw.text((x + 90, y + 25), heading, font=node_font, fill=accent)
        draw_wrapped(draw, text, (x + 28, y + 86), body_font, "#334155", card_w - 56, line_gap=6)
        if i < 5:
            if i == 2:
                x1, y1 = positions[i][0] + card_w / 2, positions[i][1] + card_h + 8
                x2, y2 = positions[i + 1][0] + card_w / 2, positions[i + 1][1] - 12
                draw.line((x1, y1, x2, y2), fill="#B8B1A5", width=5)
                draw.polygon([(x2, y2), (x2 - 12, y2 - 20), (x2 + 12, y2 - 20)], fill="#B8B1A5")
            elif i != 2:
                x1, y1 = positions[i][0] + card_w + 8, positions[i][1] + card_h / 2
                x2, y2 = positions[i + 1][0] - 12, positions[i + 1][1] + card_h / 2
                draw.line((x1, y1, x2, y2), fill="#B8B1A5", width=5)
                draw.polygon([(x2, y2), (x2 - 20, y2 - 12), (x2 - 20, y2 + 12)], fill="#B8B1A5")

    draw.rounded_rectangle((120, 872, 1480, 986), radius=24, fill="#26313A")
    draw.text((154, 898), "一句话版本", font=node_font, fill="#F8FAFC")
    draw_wrapped(
        draw,
        "读取我的日历和状态，把忙碌、恢复、低电量和一点点变好，整理成漂亮可分享的生活日历。",
        (154, 940),
        body_font,
        "#F8FAFC",
        1230,
        line_gap=8,
    )
    img.save(OUT_PNG, quality=96)


def build_doc():
    create_infographic()
    doc = Document()

    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    styles["Normal"].font.name = "Microsoft YaHei"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    styles["Normal"].font.size = Pt(10.8)
    styles["Normal"].font.color.rgb = INK
    styles["Normal"].paragraph_format.space_after = Pt(6)
    styles["Normal"].paragraph_format.line_spacing = 1.25

    p = doc.add_paragraph()
    set_para(p, before=0, after=6, line=1.1, align=WD_ALIGN_PARAGRAPH.CENTER)
    r = p.add_run("给梁姐看的需求演化简报")
    set_run_font(r, size=25, color=INK, bold=True)
    p = doc.add_paragraph()
    set_para(p, after=14, line=1.2, align=WD_ALIGN_PARAGRAPH.CENTER)
    r = p.add_run("从“番茄时钟”到“生活时间显影器”")
    set_run_font(r, size=14, color=MUTED)
    p = doc.add_paragraph()
    set_para(p, after=20, line=1.2, align=WD_ALIGN_PARAGRAPH.CENTER)
    r = p.add_run(f"整理日期：{date.today().isoformat()} | 版本：沟通后概念稿")
    set_run_font(r, size=9.5, color=MUTED)

    add_callout(
        doc,
        "核心结论",
        "这个需求最后并不是一个番茄时钟。它更像一个“生活时间显影器”：读取我已经在日历里留下的生活痕迹，听我聊当天状态，在合适的时候提醒吃饭和恢复，最后生成漂亮、可回看、可分享的日历图。",
        fill=LIGHT_TEAL,
        accent=TEAL,
    )

    doc.add_picture(str(OUT_PNG), width=Inches(6.35))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_p(doc, "上图可以直接当作给梁姐看的“一页故事”：它保留了需求被一步步修正的过程。", size=9.6, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, after=10)

    add_heading(doc, "1. 我们最初以为：想要一个不付费的番茄时钟", 1)
    add_p(doc, "起点很朴素：市面上的番茄时钟越来越爱收费、订阅、加限制，所以想做一个自己可控、安静、无广告、无订阅的工具。")
    add_bullets(doc, [
        "最初想象里，它有专注、短休息、长休息。",
        "它可以自定义 25 分钟、5 分钟、15 分钟。",
        "它可以记录今天完成了几个番茄。",
    ])
    add_callout(
        doc,
        "第一次转折",
        "很快我们发现：真正困难的不是“计时”，而是“开始”。如果开始前还要设置多久，本身就变成了一个新的门槛。",
        fill=LIGHT_BLUE,
        accent=BLUE,
    )

    add_heading(doc, "2. 第一次修正：不要逼用户先估算时间", 1)
    add_p(doc, "人在开始做一件事的时候，常常并不知道它会做多久。写作、编程、整理资料、处理情绪，都很难预估时长。所以传统倒计时并不总是合适。")
    add_bullets(doc, [
        "前台应该只有一个很轻的动作：开始做一件事。",
        "默认可以正计时，用户手动结束。",
        "时钟在后台记录，不要求用户先承诺。",
        "结束后再轻轻整理：这段算完成、热身、被打断，还是不记录。",
    ])

    add_heading(doc, "3. 第二次修正：这不是强迫专注，而是接住低电量", 1)
    add_p(doc, "继续聊下去以后，底层需求从“我要更高效”变成了“我经常低电量生活，需要有人帮我判断今天该怎么开始”。")
    add_bullets(doc, [
        "睡眠少，不应该硬开高强度任务。",
        "没吃饭，效率低可能不是懒，而是能量不足。",
        "咖啡因太晚，可能影响晚间恢复。",
        "已经连续忙很久，提醒应该是休息，而不是再坚持。",
    ])
    add_callout(
        doc,
        "产品语气",
        "它不应该像绩效工具，而应该像一个温和的生活助手：你不是不够努力，你可能只是没电了。",
        fill=LIGHT_ROSE,
        accent=ROSE,
    )

    add_heading(doc, "4. 最迷人的发现：真正的动机是回看和分享", 1)
    add_p(doc, "最后我们发现，用户愿意持续整理时间，可能不是因为计时本身，而是因为最后会得到一张漂亮的总结图、一本可以翻回去看的生活日历。")
    add_bullets(doc, [
        "看到“原来我有一天天变好”。",
        "看到“原来最近这么忙，难怪我累”。",
        "低电量日也被记录，不被当成失败。",
        "分享图带来满足感，也让整理时间这件事本身有了回报。",
    ])

    add_heading(doc, "5. 最终产品定义", 1)
    add_callout(
        doc,
        "一句话",
        "读取我的日历和状态，把忙碌、恢复、低电量和一点点变好，整理成漂亮可分享的生活日历。",
        fill=LIGHT_GOLD,
        accent=GOLD,
    )
    add_p(doc, "这个定义里，番茄时钟只是一个可选采集方式。真正核心是：已有日历数据、聊天式状态记录、温柔提醒、漂亮总结图。")

    add_heading(doc, "6. MVP 应该做什么", 1)
    add_numbered(doc, [
        "读取已有日历：先支持 Google 日历、Windows/Outlook 日历，或从 ICS 导入开始。",
        "自动归类时间：识别工作、会议、吃饭、休息、通勤、运动、社交等块。",
        "聊天式记录状态：允许用户说“今天睡得差”“中午没吃”“今天状态不错”。",
        "生活提醒：结合日历空档提醒吃饭、喝水、咖啡因、休息。",
        "生成总结图：今日卡片、最近 30 天热力图、周/月总结图。",
        "隐私分享：默认隐藏具体日程标题，只展示总结后的状态和图形。",
    ])

    add_heading(doc, "7. 不应该做什么", 1)
    add_bullets(doc, [
        "不再造一个新的日历，让用户迁移所有日程。",
        "不把设置计时作为核心入口。",
        "不把热力图做成越忙越好的排行榜。",
        "不把低电量、恢复日、空白日设计成失败。",
        "不做复杂表单式健康打卡。",
    ])

    add_heading(doc, "8. 推荐技术路径", 1)
    add_p(doc, "如果这个方向成立，技术路径也要从纯 PWA 番茄钟转向“日历集成 + 总结生成”。")
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table_width(table, [1.45, 2.65, 2.4])
    set_table_borders(table)
    headers = ["阶段", "建议做法", "原因"]
    for i, h in enumerate(headers):
        cell = table.cell(0, i)
        set_cell_shading(cell, LIGHT_GRAY)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        set_para(p, after=0)
        r = p.add_run(h)
        set_run_font(r, size=10.2, color=INK, bold=True)
    rows = [
        ("概念验证", "ICS 导入或连接一个主日历；先做 30 天热力图和今日卡片", "最快验证“漂亮回看图”是否有吸引力"),
        ("第一版", "接入 Google Calendar 或 Microsoft Graph；增加聊天式状态记录", "让数据来自用户已有生活，而不是新建系统"),
        ("第二版", "轻量后端处理授权、定时总结和提醒", "自动提醒吃饭、咖啡因和恢复通常需要后台能力"),
    ]
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_margins(cells[i])
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cells[i].paragraphs[0]
            set_para(p, after=0, line=1.2)
            r = p.add_run(value)
            set_run_font(r, size=9.5, color=INK)

    add_heading(doc, "9. 给梁姐看的判断", 1)
    add_p(doc, "这段需求的迷人之处在于，它从一个工具诉求慢慢长成了一个情绪和生活诉求：用户不是想被管理，而是想被看见。")
    add_bullets(doc, [
        "它的商业或产品吸引力不在计时器，而在“漂亮证据”。",
        "它的长期粘性不在打卡压力，而在翻日历的满足感。",
        "它的差异化不在效率，而在把低电量生活也温柔地归档。",
    ])
    add_callout(
        doc,
        "最短版本",
        "我愿意认真整理我的日历，不是因为我爱管理时间，而是因为最后它会变成一本漂亮的、证明我确实生活过的日历。",
        fill=LIGHT_TEAL,
        accent=TEAL,
    )

    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run("生活时间显影器 | 需求演化简报")
    set_run_font(r, size=8.5, color=MUTED)

    doc.save(OUT_DOCX)


if __name__ == "__main__":
    build_doc()
    print(OUT_DOCX)
    print(OUT_PNG)
